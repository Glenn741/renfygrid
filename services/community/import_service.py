"""Carga de formularios de las guias desde Excel o archivo plano (Track E,
E1, migracion 0047).

Pedido del usuario (19-oct): lo que se levanta en un taller "debe poder
cargarse post-taller mediante carga por archivo plano o Excel".

- Plantilla Excel generada desde el catalogo (nada escrito a mano aqui):
  una hoja "Operadores" y una hoja por formulario de las guias elegidas. La
  celda A1 de cada hoja lleva el codigo del formulario; asi el archivo se
  reconoce aunque alguien cambie el nombre de la hoja.
  - lista: una fila por operador e item (Respuesta con desplegable);
  - ficha: una fila por operador y campo (Valor, ¿Estimado?, nota);
  - tabla: una fila por fila de la tabla, con la columna Operador.
- Archivo plano (CSV, separado por coma o punto y coma): una fila por dato,
  columnas operador, formulario, campo, fila, valor, estimado, observacion,
  accion, responsable, fecha_limite.
- Siempre se puede ver antes de grabar (`dry_run`): se hace todo dentro de
  una transaccion y se deshace, asi la vista previa dice exactamente lo que
  pasaria.
- Una lista solo se graba completa (como en el Portal); una ficha suma
  campos a lo que ya habia; una tabla se reemplaza con las filas del archivo.
"""

from __future__ import annotations

import csv
import io
import sys
from datetime import date, datetime
from pathlib import Path
from typing import Any

sys.path.insert(0, str(Path(__file__).resolve().parent))
sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "common"))

import psycopg  # noqa: E402

from advisory_service import (  # noqa: E402
    AdvisoryConflictError,
    _require_advisor,
    create_operator,
    find_operator_by_name,
    list_operators,
    merge_operator_record,
    save_operator_form,
    submit_operator_run,
)
from form_engine import InvalidFormError, coerce_value  # noqa: E402
from pack_engine import InvalidAnswersError, InvalidRecordError, item_scale  # noqa: E402
from pack_service import TemplateNotAvailableError, active_pack_ids, list_checklist_templates  # noqa: E402

OPERATORS_SHEET = "Operadores"
HEADER_ROW = 4
# Las listas que no son de la operacion de un acueducto no van en la
# plantilla: el cuestionario CAP es por participante, no por operador.
SKIP_KINDS = ("questionnaire",)
YES = "Sí"
NO = "No"

LIST_HEADERS = ["Operador", "Código del ítem", "Ítem", "Respuesta", "Observación", "Acción acordada", "Responsable", "Fecha límite"]
RECORD_HEADERS = ["Operador", "Código del campo", "Campo", "Valor", "¿Dato estimado? (Sí/No)", "Dónde verificarlo / quién puede ayudar"]
OP_HEADERS = ["Operador", "Comunidad o sector", "Persona de contacto", "Teléfono", "Correo"]


class ImportFileError(ValueError):
    """El archivo no se puede leer como plantilla de las guias."""


def _templates(conn: psycopg.Connection, tenant_id: str, stage_codes: list[str] | None) -> list[dict]:
    packs = active_pack_ids(conn, tenant_id)
    with conn.cursor() as cur:
        cur.execute("SELECT code, sort_order FROM process_stage WHERE pack_id = ANY(%s) AND audience = 'operator'", (packs,))
        order = dict(cur.fetchall())
    out = [t for t in list_checklist_templates(conn, tenant_id)
           if t["stage_code"] in order and t["kind"] not in SKIP_KINDS and (not stage_codes or t["stage_code"] in stage_codes)]
    return sorted(out, key=lambda t: order.get(t["stage_code"], 99))


def _sheet_name(template: dict, used: set[str]) -> str:
    base = f"{template['stage_code']} {template['title']}"
    for bad in "[]:*?/\\":
        base = base.replace(bad, " ")
    name = base[:31].strip()
    n = 2
    while name.lower() in used:
        suffix = f" ({n})"
        name = base[:31 - len(suffix)].strip() + suffix
        n += 1
    used.add(name.lower())
    return name


# ── Plantilla Excel ───────────────────────────────────────────────────

def build_template(conn: psycopg.Connection, tenant_id: str, stage_codes: list[str] | None = None,
                   operator_ids: list[str] | None = None) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    _require_advisor(conn, tenant_id)
    operators = [op for op in list_operators(conn, tenant_id) if op["status"] == "active"]
    if operator_ids:
        wanted = set(operator_ids)
        operators = [op for op in operators if op["operator_id"] in wanted]
    names = [op["name"] for op in operators] or [""]
    templates = _templates(conn, tenant_id, stage_codes)

    bold = Font(bold=True)
    head_fill = PatternFill("solid", fgColor="E0E7FF")
    wrap = Alignment(wrap_text=True, vertical="top")

    wb = Workbook()
    info = wb.active
    info.title = "Instrucciones"
    lines = [
        "Plantilla de carga de los formularios de las guías",
        "",
        "1. En la hoja Operadores, revise los operadores. Puede agregar operadores nuevos al final (solo el nombre es obligatorio).",
        "2. Cada hoja es un formulario de una guía. La celda A1 tiene su código: no la cambie.",
        "3. Listas: elija la respuesta de cada ítem en la columna Respuesta. Una lista se graba solo si tiene todos sus ítems respondidos.",
        "4. Fichas: escriba el valor de cada campo. Si el dato no es exacto, márquelo como estimado y anote dónde verificarlo.",
        "   Ubicación: latitud y longitud separadas por coma (p. ej. -2.8974, -78.7756).",
        "5. Tablas: una fila por cada fila de la tabla, con el nombre del operador en la primera columna.",
        "6. Puede borrar las filas u hojas que no use. Al cargar el archivo, la plataforma muestra lo que va a grabar antes de hacerlo.",
    ]
    for i, line in enumerate(lines, start=1):
        info.cell(row=i, column=1, value=line).font = bold if i == 1 else Font()
    info.column_dimensions["A"].width = 120

    ws = wb.create_sheet(OPERATORS_SHEET)
    for c, h in enumerate(OP_HEADERS, start=1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.font, cell.fill = bold, head_fill
        ws.column_dimensions[cell.column_letter].width = 28
    for r, op in enumerate(operators, start=2):
        for c, key in enumerate(("name", "community", "contact_name", "contact_phone", "contact_email"), start=1):
            ws.cell(row=r, column=c, value=op.get(key))

    used = {"instrucciones", OPERATORS_SHEET.lower()}
    for t in templates:
        ws = wb.create_sheet(_sheet_name(t, used))
        ws.cell(row=1, column=1, value=t["id"]).font = Font(bold=True, color="64748B")
        ws.cell(row=1, column=2, value=t["title"]).font = Font(bold=True, size=12)
        ws.cell(row=2, column=2, value=t["purpose"]).alignment = wrap
        ws.row_dimensions[2].height = 45
        spec = t["form_spec"] or {}
        if t["kind"] == "record":
            headers = RECORD_HEADERS
        elif t["kind"] == "table":
            headers = ["Operador"] + [c["label"] + (" *" if c.get("required") else "") for c in spec.get("columns") or []]
        else:
            headers = LIST_HEADERS
        for c, h in enumerate(headers, start=1):
            cell = ws.cell(row=HEADER_ROW, column=c, value=h)
            cell.font, cell.fill, cell.alignment = bold, head_fill, wrap
            ws.column_dimensions[cell.column_letter].width = 22
        ws.column_dimensions["A"].width = 28
        ws.freeze_panes = ws.cell(row=HEADER_ROW + 1, column=2)
        row = HEADER_ROW + 1

        if t["kind"] == "record":
            ws.column_dimensions["C"].width = 45
            ws.column_dimensions["D"].width = 30
            ws.column_dimensions["F"].width = 34
            yes_no = DataValidation(type="list", formula1=f'"{YES},{NO}"', allow_blank=True)
            ws.add_data_validation(yes_no)
            for name in names:
                for f in spec.get("fields") or []:
                    ws.cell(row=row, column=1, value=name)
                    ws.cell(row=row, column=2, value=f["key"])
                    ws.cell(row=row, column=3, value=f["label"] + (" *" if f.get("required") else "")).alignment = wrap
                    ws.cell(row=row, column=6, value=f.get("verify"))
                    if f["type"] == "choice" and f.get("options"):
                        dv = DataValidation(type="list", formula1='"' + ",".join(o["label"].replace(",", " ") for o in f["options"]) + '"',
                                            allow_blank=True)
                        ws.add_data_validation(dv)
                        dv.add(ws.cell(row=row, column=4))
                    yes_no.add(ws.cell(row=row, column=5))
                    row += 1
        elif t["kind"] == "table":
            columns = spec.get("columns") or []
            for idx, col in enumerate(columns, start=2):
                if col["type"] == "choice" and col.get("options"):
                    dv = DataValidation(type="list", formula1='"' + ",".join(o["label"].replace(",", " ") for o in col["options"]) + '"',
                                        allow_blank=True)
                    ws.add_data_validation(dv)
                    dv.add(f"{ws.cell(row=HEADER_ROW + 1, column=idx).column_letter}{HEADER_ROW + 1}:"
                           f"{ws.cell(row=HEADER_ROW + 1, column=idx).column_letter}2000")
            preset = spec.get("preset_rows") or []
            blank = max(spec.get("suggested_rows") or 3, spec.get("min_rows") or 1)
            for name in names:
                rows = preset or [{}] * blank
                for r in rows:
                    ws.cell(row=row, column=1, value=name)
                    for idx, col in enumerate(columns, start=2):
                        if col["key"] in r:
                            ws.cell(row=row, column=idx, value=r[col["key"]])
                    row += 1
        else:
            ws.column_dimensions["C"].width = 60
            labels: list[str] = []
            for item in t["items"]:
                for e in item_scale(t, item):
                    if e["label"] not in labels:
                        labels.append(e["label"])
            dv = DataValidation(type="list", formula1='"' + ",".join(lab.replace(",", " ") for lab in labels) + '"', allow_blank=True)
            ws.add_data_validation(dv)
            dv.add(f"D{HEADER_ROW + 1}:D5000")
            for name in names:
                for item in t["items"]:
                    ws.cell(row=row, column=1, value=name)
                    ws.cell(row=row, column=2, value=item["key"])
                    text = (f"[{item['group']}] " if item.get("group") else "") + item["text"]
                    ws.cell(row=row, column=3, value=text).alignment = wrap
                    row += 1
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ── Lectura del archivo ───────────────────────────────────────────────

def _cell(value: Any) -> Any:
    if isinstance(value, str):
        value = value.strip()
        return value or None
    return value


def _blank_parsed() -> dict:
    return {"operators": [], "forms": {}, "errors": []}


def _put(parsed: dict, op_name: str, template_id: str, kind: str) -> dict:
    by_op = parsed["forms"].setdefault(op_name.strip(), {})
    if template_id not in by_op:
        by_op[template_id] = {"kind": kind, "answers": {}, "fields": {}, "rows": {}}
    return by_op[template_id]


def parse_xlsx(content: bytes, templates: dict[str, dict]) -> dict:
    from openpyxl import load_workbook
    try:
        wb = load_workbook(io.BytesIO(content), data_only=True, read_only=True)
    except Exception as exc:  # zip/xml invalido
        raise ImportFileError(f"No se pudo leer el Excel: {exc}") from None
    parsed = _blank_parsed()
    for ws in wb.worksheets:
        rows = list(ws.iter_rows(values_only=True))
        if ws.title.strip().lower() == OPERATORS_SHEET.lower():
            for n, r in enumerate(rows[1:], start=2):
                r = [_cell(v) for v in (list(r) + [None] * 5)[:5]]
                if r[0]:
                    parsed["operators"].append({"name": str(r[0]), "community": r[1], "contact_name": r[2],
                                                "contact_phone": None if r[3] is None else str(r[3]), "contact_email": r[4]})
            continue
        if not rows or not rows[0] or _cell(rows[0][0]) not in templates:
            continue
        t = templates[_cell(rows[0][0])]
        data_rows = rows[HEADER_ROW:]
        for n, r in enumerate(data_rows, start=HEADER_ROW + 1):
            r = [_cell(v) for v in r]
            if not r:
                continue
            op = r[0]
            if t["kind"] == "table":
                values = r[1:]
                if not any(v is not None for v in values):
                    continue
                if not op:
                    parsed["errors"].append({"sheet": ws.title, "row": n, "message": "Falta el nombre del operador"})
                    continue
                cols = (t["form_spec"] or {}).get("columns") or []
                row_data = {c["key"]: values[i] for i, c in enumerate(cols) if i < len(values) and values[i] is not None}
                # Una fila que solo trae lo que la plantilla ya traia escrito (p. ej. los temas del 7F) no es un dato.
                if row_data in ((t["form_spec"] or {}).get("preset_rows") or []):
                    continue
                slot = _put(parsed, str(op), t["id"], "table")
                slot["rows"][len(slot["rows"]) + 1] = row_data
                continue
            key = r[1] if len(r) > 1 else None
            if t["kind"] == "record":
                value = r[3] if len(r) > 3 else None
                estimated = r[4] if len(r) > 4 else None
                note = r[5] if len(r) > 5 else None
                if value is None and estimated is None:
                    continue
                if not op:
                    parsed["errors"].append({"sheet": ws.title, "row": n, "message": "Falta el nombre del operador"})
                    continue
                # La nota de la plantilla viene pre-llenada con "donde verificar": solo cuenta si hay valor o es distinta.
                field = next((f for f in (t["form_spec"] or {}).get("fields") or [] if f["key"] == key), None)
                if field is not None and note == field.get("verify"):
                    note = None
                _put(parsed, str(op), t["id"], "record")["fields"][str(key)] = {"value": value, "estimated": estimated, "note": note}
            else:
                answer = r[3] if len(r) > 3 else None
                if answer is None:
                    continue
                if not op:
                    parsed["errors"].append({"sheet": ws.title, "row": n, "message": "Falta el nombre del operador"})
                    continue
                _put(parsed, str(op), t["id"], "list")["answers"][str(key)] = {
                    "answer": answer, "observation": r[4] if len(r) > 4 else None, "action": r[5] if len(r) > 5 else None,
                    "responsible": r[6] if len(r) > 6 else None, "due_date": r[7] if len(r) > 7 else None, "row": n,
                    "sheet": ws.title}
    return parsed


CSV_COLUMNS = ("operador", "formulario", "campo", "fila", "valor", "estimado", "observacion", "accion", "responsable", "fecha_limite")


def parse_csv(content: bytes, templates: dict[str, dict]) -> dict:
    text = None
    for encoding in ("utf-8-sig", "latin-1"):
        try:
            text = content.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    if text is None:
        raise ImportFileError("No se pudo leer el archivo como texto")
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    reader = csv.DictReader(io.StringIO(text), dialect=dialect)
    header = [h.strip().lower() for h in reader.fieldnames or []]
    missing = [c for c in ("operador", "formulario", "campo", "valor") if c not in header]
    if missing:
        raise ImportFileError(f"Al archivo plano le faltan las columnas: {', '.join(missing)} (columnas: {', '.join(CSV_COLUMNS)})")
    parsed = _blank_parsed()
    for n, raw in enumerate(reader, start=2):
        r = {k.strip().lower(): _cell(v) for k, v in raw.items() if k}
        op, tid, key = r.get("operador"), r.get("formulario"), r.get("campo")
        if not any(r.values()):
            continue
        if not op or not tid or not key:
            parsed["errors"].append({"sheet": "archivo plano", "row": n, "message": "Faltan operador, formulario o campo"})
            continue
        t = templates.get(tid)
        if t is None:
            parsed["errors"].append({"sheet": "archivo plano", "row": n, "message": f"Formulario desconocido: {tid}"})
            continue
        if t["kind"] == "table":
            fila = r.get("fila") or "1"
            _put(parsed, op, tid, "table")["rows"].setdefault(str(fila), {})[key] = r.get("valor")
        elif t["kind"] == "record":
            _put(parsed, op, tid, "record")["fields"][key] = {"value": r.get("valor"), "estimated": r.get("estimado"),
                                                              "note": r.get("observacion")}
        else:
            _put(parsed, op, tid, "list")["answers"][key] = {
                "answer": r.get("valor"), "observation": r.get("observacion"), "action": r.get("accion"),
                "responsible": r.get("responsable"), "due_date": r.get("fecha_limite"), "row": n, "sheet": "archivo plano"}
    return parsed


# ── Grabar ────────────────────────────────────────────────────────────

def _list_answers(t: dict, raw: dict[str, dict]) -> list[dict]:
    items = {i["key"]: i for i in t["items"]}
    unknown = sorted(set(raw) - set(items))
    if unknown:
        raise InvalidAnswersError(f"Ítems que no existen en la lista: {unknown}")
    answers = []
    for key, item in items.items():
        a = raw.get(key)
        if a is None:
            continue
        text = str(a["answer"]).strip().lower()
        entry = next((e for e in item_scale(t, item) if text in (e["code"].lower(), e["label"].lower())
                      or e["label"].lower().startswith(text + " ")), None)
        if entry is None:
            labels = ", ".join(e["label"] for e in item_scale(t, item))
            raise InvalidAnswersError(f"«{item['text'][:60]}»: {a['answer']!r} no es una respuesta válida ({labels})")
        due = a.get("due_date")
        if due is not None:
            due = coerce_value({"key": "due", "label": "Fecha límite", "type": "date"}, due)
        answers.append({"item_key": key, "answer_code": entry["code"], "observation": _str(a.get("observation")),
                        "action": _str(a.get("action")), "responsible": _str(a.get("responsible")), "due_date": due})
    missing = len(items) - len(answers)
    if missing:
        raise InvalidAnswersError(f"Faltan {missing} de {len(items)} ítems por responder: la lista se graba solo completa")
    return answers


def _str(value: Any) -> str | None:
    if value is None:
        return None
    if isinstance(value, (datetime, date)):
        return value.isoformat()[:10]
    return str(value).strip() or None


def import_file(conn: psycopg.Connection, tenant_id: str, actor: str, filename: str, content: bytes,
                dry_run: bool = True, workshop_id: str | None = None, create_missing: bool = True) -> dict:
    _require_advisor(conn, tenant_id)
    templates = {t["id"]: t for t in _templates(conn, tenant_id, None)}
    name = (filename or "").lower()
    if name.endswith(".xlsx") or content[:2] == b"PK":
        parsed = parse_xlsx(content, templates)
    elif name.endswith((".csv", ".txt")):
        parsed = parse_csv(content, templates)
    else:
        raise ImportFileError("Use la plantilla Excel (.xlsx) o un archivo plano (.csv)")

    result: dict[str, Any] = {"dry_run": dry_run, "operators_created": [], "lists_saved": 0, "forms_saved": 0,
                              "by_operator": [], "errors": list(parsed["errors"])}
    declared = {o["name"].strip().lower(): o for o in parsed["operators"]}
    names = {o["name"].strip(): None for o in parsed["operators"]}
    names.update({n: None for n in parsed["forms"]})

    try:
        with conn.transaction():
            for op_name in names:
                existing = find_operator_by_name(conn, tenant_id, op_name)
                if existing is None:
                    if not create_missing and op_name.lower() not in declared:
                        result["errors"].append({"sheet": "", "row": None,
                                                 "message": f"{op_name}: no está registrado (agréguelo en la hoja Operadores)"})
                        continue
                    extra = declared.get(op_name.lower(), {})
                    existing = create_operator(conn, tenant_id, actor, op_name, community=extra.get("community"),
                                               contact_name=extra.get("contact_name"), contact_phone=extra.get("contact_phone"),
                                               contact_email=extra.get("contact_email"), origin="import", workshop_id=workshop_id)
                    result["operators_created"].append(op_name)
                summary = {"operator": op_name, "lists": [], "forms": []}
                for tid, slot in (parsed["forms"].get(op_name) or {}).items():
                    t = templates[tid]
                    try:
                        if slot["kind"] == "list":
                            answers = _list_answers(t, slot["answers"])
                            submit_operator_run(conn, tenant_id, existing["operator_id"], tid, answers, actor,
                                                notes=f"Cargado desde {filename}", workshop_id=workshop_id)
                            result["lists_saved"] += 1
                            summary["lists"].append(t["title"])
                        elif slot["kind"] == "record":
                            saved = merge_operator_record(conn, tenant_id, existing["operator_id"], tid, {"fields": slot["fields"]},
                                                          actor, workshop_id=workshop_id)
                            result["forms_saved"] += 1
                            summary["forms"].append(f"{t['title']} ({'completa' if saved['progress']['complete'] else 'en borrador'})")
                        else:
                            rows = [slot["rows"][k] for k in slot["rows"]]
                            saved = save_operator_form(conn, tenant_id, existing["operator_id"], tid, {"rows": rows}, actor,
                                                       workshop_id=workshop_id, source="import")
                            result["forms_saved"] += 1
                            summary["forms"].append(f"{t['title']} ({len(saved['data']['rows'])} filas)")
                    except (InvalidAnswersError, InvalidFormError, InvalidRecordError, AdvisoryConflictError,
                            TemplateNotAvailableError) as exc:
                        result["errors"].append({"sheet": t["title"], "row": None, "message": f"{op_name}: {exc}"})
                result["by_operator"].append(summary)
            if dry_run:
                raise psycopg.Rollback()
    except psycopg.Rollback:
        pass
    return result
