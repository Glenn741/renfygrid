"""Carga por archivo de fronteras y equipos (Track F, Q2, 2026-10-20).

Un comercializador llega con sus fronteras y sus medidores en hojas de
calculo. La plantilla (`build_template`) tiene dos hojas:

- Fronteras: una fila por frontera (se crea, o se actualiza si el codigo ya
  existe).
- Equipos: una fila por equipo (medidor, CT, PT o modem). Si trae frontera y
  posicion se instala; si trae fecha de ultima calibracion o de prueba de
  rutina queda en su bitacora con el certificado.

Los codigos validos (tipos de punto, tipos de equipo, posiciones, estados)
salen del catalogo del paquete y se escriben en la hoja "Códigos" de la
plantilla. Con `dry_run` (por defecto) dice lo que grabaria sin grabar nada.
Cada fila corre en su propio punto de guardado: una fila con error no tumba
las demas y se informa con su numero.
"""

from __future__ import annotations

import csv
import io
import sys
from datetime import date, datetime
from pathlib import Path
from typing import Any

sys.path.insert(0, str(Path(__file__).resolve().parent))
sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "common"))
sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "community"))

import openpyxl  # noqa: E402
import psycopg  # noqa: E402
from openpyxl.styles import Font, PatternFill  # noqa: E402

import metering_asset_service as svc  # noqa: E402
from form_engine import InvalidFormError, excel_serial_to_date  # noqa: E402
from renmeter_common.db import tenant_scope  # noqa: E402

POINT_SHEET = "Fronteras"
EQUIPMENT_SHEET = "Equipos"
CODES_SHEET = "Códigos"

# (clave, encabezado, obligatorio). Las claves son las columnas del modelo.
POINT_COLUMNS = [
    ("code", "Código de la frontera (SIC)", True), ("name", "Nombre de la frontera", True),
    ("customer_name", "Usuario (cliente)", True), ("address", "Dirección", False), ("city", "Ciudad o municipio", False),
    ("region", "Departamento", False), ("lat", "Latitud", False), ("lon", "Longitud", False),
    ("point_class", "Tipo de punto", False), ("voltage_kv", "Tensión de servicio (kV)", False),
    ("capacity_kva", "Capacidad instalada (kVA)", False), ("consumption_mwh_month", "Consumo mensual (MWh)", False),
    ("status", "Estado", False), ("registered_on", "Fecha de registro ante el ASIC", False),
    ("verified_on", "Fecha de verificación inicial", False), ("export_code", "Código SIC de exportación", False),
    ("niu", "NIU", False),
]
EQUIPMENT_COLUMNS = [
    ("class_code", "Tipo de equipo", True), ("serial", "Número de serie", True), ("brand", "Marca", False),
    ("model", "Modelo", False), ("manufacturer", "Fabricante", False), ("year", "Año de fabricación", False),
    ("ownership", "Propiedad", False), ("point_code", "Frontera donde está instalado", False),
    ("position", "Posición", False), ("installed_on", "Fecha de instalación", False),
    ("last_calibration_on", "Fecha de la última calibración", False),
    ("calibration_certificate", "Certificado de calibración", False),
    ("calibration_laboratory", "Laboratorio de calibración", False),
    ("last_routine_test_on", "Fecha de la última prueba de rutina", False),
    ("routine_test_certificate", "Certificado de prueba de rutina", False),
]


class AssetImportError(ValueError):
    """El archivo no se puede leer como plantilla de fronteras y equipos."""


# ── Plantilla ─────────────────────────────────────────────────────────

def build_template(conn: psycopg.Connection, tenant_id: str) -> bytes:
    cat = svc.catalog(conn, tenant_id)
    wb = openpyxl.Workbook()
    bold = Font(bold=True)
    required_fill = PatternFill("solid", fgColor="FDE68A")
    for index, (title, columns) in enumerate(((POINT_SHEET, POINT_COLUMNS), (EQUIPMENT_SHEET, EQUIPMENT_COLUMNS))):
        ws = wb.active if index == 0 else wb.create_sheet()
        ws.title = title
        for col, (_, header, required) in enumerate(columns, start=1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = bold
            if required:
                cell.fill = required_fill
            ws.column_dimensions[cell.column_letter].width = max(14, len(header) + 2)
        ws.freeze_panes = "A2"
    codes = wb.create_sheet(CODES_SHEET)
    codes.append(["Lista", "Código", "Significado"])
    for c in codes[1]:
        c.font = bold
    for pc in cat["point_classes"]:
        codes.append(["Tipo de punto", pc["code"], f"{pc['label']}: consumo ≥ {pc['min_mwh_month']:g} MWh-mes o capacidad ≥ {pc['min_mva']:g} MVA"])
    for ec in cat["equipment_classes"]:
        codes.append(["Tipo de equipo", ec["code"], ec["label"]])
    for p in cat["positions"]:
        codes.append(["Posición", p["code"], p["label"]])
    for code, label in _labels(conn, "metering_point.status"):
        codes.append(["Estado", code, label])
    for code, label in _labels(conn, "equipment.ownership"):
        codes.append(["Propiedad", code, label])
    codes.column_dimensions["A"].width = 16
    codes.column_dimensions["B"].width = 16
    codes.column_dimensions["C"].width = 80
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def _labels(conn: psycopg.Connection, domain: str) -> list[tuple[str, str]]:
    with conn.cursor() as cur:
        cur.execute("SELECT code, label FROM code_label WHERE domain = %s ORDER BY sort_order", (domain,))
        return [(r[0], r[1]) for r in cur.fetchall()]


# ── Lectura ───────────────────────────────────────────────────────────

def _header_map(headers: list[Any], columns: list[tuple[str, str, bool]]) -> dict[int, str]:
    lookup = {}
    for key, header, _ in columns:
        lookup[key.lower()] = key
        lookup[header.lower()] = key
    out = {}
    for i, h in enumerate(headers):
        key = lookup.get(str(h or "").strip().lower())
        if key:
            out[i] = key
    return out


def _rows(values: list[list[Any]], columns: list[tuple[str, str, bool]]) -> list[tuple[int, dict]]:
    if not values:
        return []
    mapping = _header_map(values[0], columns)
    missing = [header for key, header, required in columns if required and key not in mapping.values()]
    if missing:
        raise AssetImportError("Faltan columnas obligatorias: " + ", ".join(missing))
    out = []
    for n, row in enumerate(values[1:], start=2):
        record = {}
        for i, key in mapping.items():
            value = row[i] if i < len(row) else None
            if isinstance(value, str):
                value = value.strip() or None
            if isinstance(value, datetime):
                value = value.date()
            if value is not None:
                record[key] = value
        if record:
            out.append((n, record))
    return out


def parse_file(filename: str, content: bytes) -> dict[str, list[tuple[int, dict]]]:
    name = filename.lower()
    if name.endswith((".xlsx", ".xlsm")):
        try:
            wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        except Exception as exc:  # archivo dañado o que no es Excel
            raise AssetImportError("No se pudo leer el Excel; use la plantilla descargada") from exc
        out: dict[str, list[tuple[int, dict]]] = {"points": [], "equipment": []}
        if POINT_SHEET in wb.sheetnames:
            out["points"] = _rows([list(r) for r in wb[POINT_SHEET].iter_rows(values_only=True)], POINT_COLUMNS)
        if EQUIPMENT_SHEET in wb.sheetnames:
            out["equipment"] = _rows([list(r) for r in wb[EQUIPMENT_SHEET].iter_rows(values_only=True)], EQUIPMENT_COLUMNS)
        if not out["points"] and not out["equipment"]:
            raise AssetImportError(f"El Excel no tiene filas en las hojas «{POINT_SHEET}» ni «{EQUIPMENT_SHEET}»")
        return out
    if name.endswith((".csv", ".txt")):
        text = content.decode("utf-8-sig", errors="replace")
        dialect = csv.Sniffer().sniff(text[:2048], delimiters=",;\t")
        values = [list(r) for r in csv.reader(io.StringIO(text), dialect)]
        headers = {str(h).strip().lower() for h in (values[0] if values else [])}
        is_equipment = bool(headers & {"serial", "número de serie"})
        columns = EQUIPMENT_COLUMNS if is_equipment else POINT_COLUMNS
        return {"points": [] if is_equipment else _rows(values, columns), "equipment": _rows(values, columns) if is_equipment else []}
    raise AssetImportError("Formato no soportado: use la plantilla Excel (.xlsx) o un archivo plano (.csv)")


def _date(value: Any, label: str) -> date | None:
    if value is None:
        return None
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)):
        return excel_serial_to_date(value)
    text = str(value).strip()
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y"):
        try:
            return datetime.strptime(text[:10], fmt).date()
        except ValueError:
            continue
    raise InvalidFormError(f"{label}: fecha inválida ({text!r}); use AAAA-MM-DD o DD/MM/AAAA")


def _code(value: Any, options: list[dict], what: str, key: str = "code") -> str:
    text = str(value).strip().lower()
    for o in options:
        if text in (str(o[key]).lower(), str(o.get("label", "")).lower()):
            return o[key]
    raise InvalidFormError(f"{what} desconocido: {value!r} (vea la hoja «{CODES_SHEET}»)")


# ── Carga ─────────────────────────────────────────────────────────────

def import_file(conn: psycopg.Connection, tenant_id: str, imported_by: str, filename: str, content: bytes,
                dry_run: bool = True) -> dict:
    parsed = parse_file(filename, content)
    ctx = svc._context(conn, tenant_id)
    status_labels = [{"code": c, "label": lbl} for c, lbl in _labels(conn, "metering_point.status")]
    ownership_labels = [{"code": c, "label": lbl} for c, lbl in _labels(conn, "equipment.ownership")]
    classes = list(ctx["equipment_classes"].values())
    positions = list(ctx["positions"].values())
    point_classes = [{"code": c["code"], "label": c["label"]} for c in ctx["classes"]]
    result = {"dry_run": dry_run, "points_created": 0, "points_updated": 0, "equipment_created": 0, "equipment_existing": 0,
              "installed": 0, "calibrations": 0, "routine_tests": 0, "errors": []}

    def fail(sheet: str, row: int, exc: Exception) -> None:
        result["errors"].append({"sheet": sheet, "row": row, "message": str(exc)})

    try:
        with conn.transaction():
            with tenant_scope(conn, tenant_id):
                with conn.cursor() as cur:
                    for n, rec in parsed["points"]:
                        try:
                            with conn.transaction():
                                if "point_class" in rec:
                                    rec["point_class"] = _code(rec["point_class"], point_classes, "Tipo de punto")
                                if "status" in rec:
                                    rec["status"] = _code(rec["status"], status_labels, "Estado")
                                for col in ("registered_on", "verified_on"):
                                    if col in rec:
                                        rec[col] = _date(rec[col], col)
                                cur.execute("SELECT id::text FROM metering_point WHERE lower(btrim(code)) = lower(btrim(%s))",
                                            (str(rec.get("code", "")),))
                                existing = cur.fetchone()
                                if existing:
                                    clean = svc._clean_point(ctx, rec, partial=True)
                                    if "point_class" in clean:
                                        clean["pack_id"] = ctx["pack_id"] if clean["point_class"] else None
                                    sets = ", ".join(f"{c} = %s" for c in clean)
                                    cur.execute(f"UPDATE metering_point SET {sets}, updated_at = now() WHERE id = %s",
                                                [*clean.values(), existing[0]])
                                    result["points_updated"] += 1
                                else:
                                    svc._create_point(cur, tenant_id, ctx, rec, imported_by)
                                    result["points_created"] += 1
                        except (InvalidFormError, svc.AssetConflictError, svc.AssetNotFoundError, psycopg.DataError) as exc:
                            fail(POINT_SHEET, n, exc)
                    for n, rec in parsed["equipment"]:
                        # Los contadores se suman solo si la fila entera quedo grabada:
                        # una fila con error se deshace completa en su punto de guardado.
                        counts = {k: 0 for k in ("equipment_created", "equipment_existing", "installed", "calibrations",
                                                 "routine_tests")}
                        try:
                            with conn.transaction():
                                _import_equipment(cur, tenant_id, ctx, rec, imported_by, classes, positions,
                                                  ownership_labels, counts)
                            for k, v in counts.items():
                                result[k] += v
                        except (InvalidFormError, svc.AssetConflictError, svc.AssetNotFoundError, psycopg.DataError) as exc:
                            fail(EQUIPMENT_SHEET, n, exc)
            if dry_run:
                raise psycopg.Rollback()
    except psycopg.Rollback:
        pass
    return result


def _import_equipment(cur: psycopg.Cursor, tenant_id: str, ctx: dict, rec: dict, by: str, classes: list[dict],
                      positions: list[dict], ownership_labels: list[dict], result: dict) -> None:
    rec["class_code"] = _code(rec["class_code"], classes, "Tipo de equipo")
    rec["serial"] = str(rec["serial"]).strip()
    if "ownership" in rec:
        rec["ownership"] = _code(rec["ownership"], ownership_labels, "Propiedad")
    data = {k: rec[k] for k in ("class_code", "serial", "brand", "model", "manufacturer", "year", "ownership") if k in rec}
    if isinstance(data.get("serial"), float):
        data["serial"] = f"{data['serial']:.0f}"
    equipment_id = svc.find_equipment(cur, data["class_code"], data.get("brand"), data["serial"])
    if equipment_id:
        result["equipment_existing"] += 1
    else:
        equipment_id = svc._create_equipment(cur, tenant_id, ctx, data, by)
        result["equipment_created"] += 1
    installed_on = _date(rec.get("installed_on"), "Fecha de instalación")
    if rec.get("point_code"):
        cur.execute("SELECT id::text FROM metering_point WHERE lower(btrim(code)) = lower(btrim(%s))", (str(rec["point_code"]),))
        point = cur.fetchone()
        if point is None:
            raise svc.AssetNotFoundError(f"No existe la frontera {rec['point_code']!r} (cárguela en la hoja «{POINT_SHEET}»)")
        if not rec.get("position"):
            raise InvalidFormError("Indique la posición en la que está instalado")
        position = _code(rec["position"], positions, "Posición")
        cur.execute("SELECT point_id::text, position_code FROM equipment_installation WHERE equipment_id = %s AND removed_on IS NULL",
                    (equipment_id,))
        current = cur.fetchone()
        if current is None:
            if installed_on is None:
                raise InvalidFormError("Indique la fecha de instalación")
            svc._install(cur, tenant_id, ctx, equipment_id, point[0], position, installed_on, by, None, "import", None)
            result["installed"] += 1
        elif current != (point[0], position):
            raise svc.AssetConflictError(f"El equipo {data['serial']} ya está instalado en otra frontera o posición")
    for activity, date_key, cert_key, lab_key, counter in (
            ("calibration", "last_calibration_on", "calibration_certificate", "calibration_laboratory", "calibrations"),
            ("routine_test", "last_routine_test_on", "routine_test_certificate", None, "routine_tests")):
        day = _date(rec.get(date_key), date_key)
        if day is None:
            continue
        cert = str(rec.get(cert_key) or "").strip()
        if not cert:
            raise InvalidFormError(f"Falta el número del certificado de {'calibración' if activity == 'calibration' else 'prueba de rutina'}")
        cur.execute("SELECT 1 FROM metering_event WHERE equipment_id = %s AND type_code = %s AND occurred_on = %s "
                    "AND data -> 'fields' -> 'certificate_number' ->> 'value' = %s", (equipment_id, activity, day, cert))
        if cur.fetchone():
            continue  # ya cargado: volver a subir el archivo no duplica la bitacora
        fields = {"certificate_number": {"value": cert}, "issued_on": {"value": day.isoformat()}, "result": {"value": "pass"}}
        if lab_key:
            lab = str(rec.get(lab_key) or "").strip()
            if not lab:
                raise InvalidFormError("Falta el laboratorio de la calibración")
            fields["laboratory"] = {"value": lab}
        svc._record(cur, tenant_id, ctx, type_code=activity, occurred_on=day, recorded_by=by, point_id=None,
                    equipment_id=equipment_id, data={"fields": fields}, notes=None, source="import", external_ref=None)
        result[counter] += 1
