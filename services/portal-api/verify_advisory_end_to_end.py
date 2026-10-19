"""Verificacion end-to-end real de Track E, Sprint E1 (modulo Guias y
asesoria de la ONG, migracion 0047) -- HTTP real (FastAPI TestClient), JWT
real, Postgres real con el rol de aplicacion (RLS activo), sin mocks.

Que prueba, en espanol llano:
  1. La plataforma crea la ONG con su administrador (y nadie mas puede); la
     ONG crea a sus asesores con nombre de usuario (sin correo).
  2. Taller del dia: el asesor lo registra con las guias que cubre, registra
     operadores (con su ubicacion y numero de usuarios) y la asistencia.
  3. El asesor llena con cada operador: listas de la G1 (sin hallazgos en la
     ONG), la ficha a medias (borrador) y luego completa, y una tabla.
  4. Carga por archivo: la plantilla Excel sale del catalogo; la vista previa
     no graba nada; la carga real crea un operador nuevo y sus formularios;
     el archivo plano (CSV) tambien.
  5. Tablero: capacitados, guias vistas por operador y el mapa con tamano,
     color y forma salidos de la ficha.
  6. Traspaso: se crea la organizacion del operador con sus datos; alli
     nacen los hallazgos; su administrador entra y ve sus guias llenas; el
     registro de la ONG queda de solo lectura.
  7. Aislamiento y permisos: otra ONG no ve nada; un prestador no usa la
     asesoria; un operador de la ONG no registra operadores.
Al final borra SOLO las organizaciones de prueba y lo que crearon.

Uso:
    python verify_advisory_end_to_end.py "postgresql://renfygrid_app:renfygrid_app_dev_only@localhost:5455/renfygrid"

Contra un servicio ya corriendo (p. ej. produccion, sin httpx en su venv):
    E2E_BASE_URL=http://10.12.200.27:8105 python verify_advisory_end_to_end.py "<dsn del servicio>"
"""

from __future__ import annotations

import base64
import io
import json as _json
import os
import sys
import urllib.error
import urllib.request
import uuid
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "common"))
sys.path.insert(0, str(Path(__file__).resolve().parent))

import psycopg  # noqa: E402
from psycopg.types.json import Json  # noqa: E402

from e2e_packs import adopt_program_packs  # noqa: E402
from renmeter_common.db import tenant_scope  # noqa: E402
from renmeter_common.user_service import create_app_user  # noqa: E402

JWT_SECRET = "e2e-e1-secret"
ORDER_SIGNING_SECRET = "e2e-e1-order-secret"
SUFFIX = uuid.uuid4().hex[:6]
CONFIG = {"session_ttl_seconds": 3600, "timezone": "America/Guayaquil", "currency": "USD", "locale": "es-EC"}


def check(condition: bool, message: str) -> None:
    if not condition:
        raise AssertionError(message)
    print(f"  OK  {message}")


class _Response:
    def __init__(self, status: int, content: bytes):
        self.status_code = status
        self.content = content
        self.text = content.decode("utf-8", errors="replace")

    def json(self):
        return _json.loads(self.content or b"null")


class _HttpClient:
    """Lo mismo que se usa de TestClient, contra un servidor real."""

    def __init__(self, base: str):
        self.base = base.rstrip("/")

    def _call(self, method: str, path: str, headers: dict | None = None, json: object = None) -> _Response:
        data = None if json is None else _json.dumps(json).encode()
        req = urllib.request.Request(self.base + path, method=method, data=data,
                                     headers={"Content-Type": "application/json", **(headers or {})})
        try:
            with urllib.request.urlopen(req, timeout=120) as r:
                return _Response(r.status, r.read())
        except urllib.error.HTTPError as e:
            return _Response(e.code, e.read())

    def get(self, path, headers=None):
        return self._call("GET", path, headers)

    def post(self, path, headers=None, json=None):
        return self._call("POST", path, headers, json)

    def put(self, path, headers=None, json=None):
        return self._call("PUT", path, headers, json)

    def patch(self, path, headers=None, json=None):
        return self._call("PATCH", path, headers, json)

    def delete(self, path, headers=None):
        return self._call("DELETE", path, headers)


def run(dsn: str) -> None:
    from openpyxl import load_workbook

    base = os.environ.get("E2E_BASE_URL")
    if base:
        client = _HttpClient(base)
    else:
        os.environ["RENFYGRID_DSN"] = dsn
        os.environ["RENFYGRID_JWT_SECRET"] = JWT_SECRET
        os.environ["RENFYGRID_ORDER_SIGNING_SECRET"] = ORDER_SIGNING_SECRET
        import main  # despues de fijar el entorno
        from fastapi.testclient import TestClient
        client = TestClient(main.app)
    created: list[str] = []

    def login(org: str, user: str, password: str) -> dict:
        r = client.post("/auth/login", json={"tenant_id": org, "email": user, "password": password})
        assert r.status_code == 200, r.text
        return {"Authorization": f"Bearer {r.json()['access_token']}"}

    with psycopg.connect(dsn, autocommit=True) as conn:
        with conn.cursor() as cur:
            for name, kind in ((f"e2e-plataforma-{SUFFIX}", "platform"), (f"e2e-junta-{SUFFIX}", "provider")):
                cur.execute("INSERT INTO tenant (name, config, kind) VALUES (%s, %s, %s) RETURNING id", (name, Json(CONFIG), kind))
                created.append(str(cur.fetchone()[0]))
        platform_id, junta_id = created
        try:
            adopt_program_packs(conn, junta_id)
            create_app_user(conn, platform_id, "root@plataforma.test", "clave-plataforma", "platform")
            create_app_user(conn, junta_id, "admin@junta.test", "clave-junta", "supervisor")
            hp = login(f"e2e-plataforma-{SUFFIX}", "root@plataforma.test", "clave-plataforma")
            hj = login(f"e2e-junta-{SUFFIX}", "admin@junta.test", "clave-junta")

            print("1. Organizaciones y usuarios")
            ong = f"e2e-apt-{SUFFIX}"
            body = {"name": ong, "kind": "advisor", "admin_username": "admin", "admin_password": "clave-admin-apt",
                    "display_name": "APT (prueba)", "packs": ["EC-ARCA", "EC-MUNICIPIOS-AZULES"]}
            check(client.post("/organizations", headers=hj, json=body).status_code == 403, "una junta no crea organizaciones")
            check(client.post("/organizations", headers=hp, json={**body, "name": "Con Espacios"}).status_code == 422,
                  "código con espacios -> 422")
            r = client.post("/organizations", headers=hp, json=body)
            check(r.status_code == 201 and r.json()["kind"] == "advisor", "la plataforma crea la ONG con su administrador")
            ong_id = r.json()["tenant_id"]
            created.append(ong_id)
            check(client.post("/organizations", headers=hp, json=body).status_code == 409, "código repetido -> 409")
            check(any(o["name"] == ong for o in client.get("/organizations", headers=hp).json()), "la plataforma lista organizaciones")
            check(client.get("/organizations", headers=hj).status_code == 403, "una junta no ve la lista")
            ha = login(ong, "admin", "clave-admin-apt")
            me = client.get("/auth/me", headers=ha).json()
            check(me["tenant_kind"] == "advisor" and me["tenant_display_name"] == "APT (prueba)", "la sesión dice que es una ONG")
            for user, role in (("asesor01", "advisor"), ("asesor02", "advisor"), ("apoyo", "operator")):
                r = client.post("/users", headers=ha, json={"email": user, "password": f"clave-{user}", "role": role})
                check(r.status_code == 201, f"la ONG crea el usuario {user} ({role}) sin correo")
            h1 = login(ong, "asesor01", "clave-asesor01")
            hop = login(ong, "apoyo", "clave-apoyo")

            print("2. Taller, operadores y asistencia")
            r = client.post("/advisory/workshops", headers=h1, json={"title": "Taller Guías 1 a 3", "held_on": "2026-10-22",
                                                                     "stage_codes": ["G1", "G2", "G3"], "place": "Gualaceo"})
            check(r.status_code == 201 and r.json()["stage_codes"] == ["G1", "G2", "G3"], "taller con las guías que cubre")
            wid = r.json()["workshop_id"]
            check(client.post("/advisory/workshops", headers=h1, json={"title": "X", "held_on": "2026-10-22",
                                                                       "stage_codes": ["G9"]}).status_code == 422, "guía inexistente -> 422")
            ops = {}
            for name, users, supply, san, lat in (("Junta San Juan", 120, "potable", "sewer", -2.89),
                                                  ("Junta El Rosal", 45, "raw", "onsite", -2.95)):
                r = client.post("/advisory/operators", headers=h1, json={
                    "name": name, "community": "Gualaceo", "workshop_id": wid,
                    "profile": {"canton": "Gualaceo", "location": {"lat": lat, "lon": -78.77}, "users": users,
                                "supply_type": supply, "sanitation": san}})
                check(r.status_code == 201, f"operador {name} registrado con ubicación y usuarios")
                ops[name] = r.json()["operator_id"]
            check(client.post("/advisory/operators", headers=h1, json={"name": " junta san juan "}).status_code == 409,
                  "nombre repetido (sin importar mayúsculas) -> 409")
            check(client.post("/advisory/operators", headers=hop, json={"name": "Otra"}).status_code == 403,
                  "el rol operador de la ONG no registra operadores")
            r = client.post(f"/advisory/workshops/{wid}/attendees", headers=h1,
                            json={"full_name": "María Pérez", "operator_id": ops["Junta San Juan"], "role_label": "Presidenta"})
            check(len(r.json()["attendees"]) == 1 and len(r.json()["operators"]) == 2, "asistencia registrada (T-02)")

            print("3. Formularios de un operador")
            sj = ops["Junta San Juan"]
            op = client.get(f"/advisory/operators/{sj}", headers=h1).json()
            g1 = next(s for s in op["route"]["stages"] if s["code"] == "G1")
            kinds = {item["template_id"]: item["kind"] for item in g1["lists"]}
            check(kinds.get("MA-G1-PROFILE") == "record" and kinds.get("MA-G1-PRIORITIES") == "table" and len(kinds) == 9,
                  "la G1 tiene sus 9 formularios (listas, fichas y tablas)")
            check([s["code"] for s in op["route"]["stages"]] == ["G1", "G2", "G3", "G4", "G5", "G6"]
                  and op["route"]["other_lists"] == [], "las herramientas del facilitador (INT) no son guías del operador")
            profile_item = next(i for i in g1["lists"] if i["template_id"] == "MA-G1-PROFILE")
            check(profile_item["status"] == "draft", "la ficha del registro queda en borrador (faltan campos obligatorios)")
            tpl = {t["id"]: t for t in client.get("/checklist-templates", headers=h1).json()}
            answers = [{"item_key": i["key"], "answer_code": "red" if i["key"] == "network" else "green"} for i in tpl["MA-G1-MAP"]["items"]]
            r = client.post(f"/advisory/operators/{sj}/checklist-runs", headers=h1,
                            json={"template_id": "MA-G1-MAP", "answers": answers, "workshop_id": wid})
            check(r.status_code == 201 and r.json()["findings_created"] == [], "mapa parlante aplicado; en la ONG no crea hallazgos")
            check(client.post(f"/advisory/operators/{sj}/checklist-runs", headers=h1,
                              json={"template_id": "MA-G1-MAP", "answers": answers[:3]}).status_code == 422, "lista incompleta -> 422")
            check(client.post(f"/advisory/operators/{sj}/checklist-runs", headers=hop,
                              json={"template_id": "MA-G1-MAP", "answers": answers}).status_code == 403, "sin permiso advisory.record -> 403")
            r = client.put(f"/advisory/operators/{sj}/forms/MA-G1-PROFILE", headers=h1, json={"data": {"fields": {
                "canton": {"value": "Gualaceo"}, "parish": {"value": "San Juan"}, "location": {"value": {"lat": -2.89, "lon": -78.77}},
                "users": {"value": 120}, "inhabitants": {"value": "480", "estimated": True, "note": "Verificar con el INEC"},
                "supply_type": {"value": "Agua potable"}, "sanitation": {"value": "sewer"},
                "economy": {"value": ["Agricultura", "trade"]}}}, "workshop_id": wid})
            check(r.status_code == 200 and r.json()["progress"]["complete"] and r.json()["progress"]["estimated"] == 1,
                  "ficha completa; habitantes marcado como estimado")
            check(r.json()["data"]["fields"]["supply_type"]["value"] == "potable"
                  and r.json()["data"]["fields"]["economy"]["value"] == ["agriculture", "trade"], "las etiquetas se guardan como códigos")
            check(client.put(f"/advisory/operators/{sj}/forms/MA-G1-PROFILE", headers=h1,
                             json={"data": {"fields": {"users": {"value": -3}}}}).status_code == 422, "número negativo -> 422")
            check(client.put(f"/advisory/operators/{sj}/forms/MA-G1-PROFILE", headers=h1,
                             json={"data": {"fields": {"inventado": {"value": 1}}}}).status_code == 422, "campo inexistente -> 422")
            r = client.put(f"/advisory/operators/{sj}/forms/MA-G1-PRIORITIES", headers=h1, json={"data": {"rows": [
                {"problem": "Fugas frecuentes en la red", "families": 70, "priority": "Alta"},
                {"problem": "Morosidad", "families": 40}]}})
            check(r.status_code == 200 and not r.json()["progress"]["complete"], "tabla con una fila sin prioridad: borrador")
            check(client.put(f"/advisory/operators/{sj}/forms/MA-G2-7D", headers=h1, json={"data": {"rows": [
                {"missing": str(i), "first_step": "x"} for i in range(4)]}}).status_code == 422, "7D admite máximo 3 filas")
            check(len(client.get(f"/advisory/operators/{sj}/checklist-runs", headers=h1).json()) == 1, "historial del operador")

            print("4. Carga por archivo")
            r = client.get("/advisory/import/template?stages=G1,G2", headers=h1)
            check(r.status_code == 200 and r.content[:2] == b"PK", "plantilla Excel generada desde el catálogo")
            wb = load_workbook(io.BytesIO(r.content))
            check("Operadores" in wb.sheetnames and len(wb.sheetnames) >= 15, f"una hoja por formulario ({len(wb.sheetnames)} hojas)")
            wb["Operadores"].append(["Junta Nueva Esperanza", "Chordeleg", "Luis", "0991234567", None])
            sheets = {ws["A1"].value: ws for ws in wb.worksheets}
            perc = sheets["MA-G1-PERCEPTION"]
            perc.delete_rows(5, perc.max_row)
            for item in tpl["MA-G1-PERCEPTION"]["items"]:
                perc.append(["Junta Nueva Esperanza", item["key"], item["text"], "Más o menos"])
            prof = sheets["MA-G1-PROFILE"]
            prof.delete_rows(5, prof.max_row)
            for key, value, est in (("canton", "Chordeleg", None), ("location", "-2.92, -78.75", None), ("users", 60, "Sí"),
                                    ("supply_type", "Agua cruda o entubada sin tratamiento", None),
                                    ("sanitation", "Saneamiento en sitio (pozo séptico  letrina  biodigestor)", None)):
                prof.append(["Junta Nueva Esperanza", key, key, value, est])
            roles = sheets["MA-G2-7A"]
            roles.delete_rows(5, roles.max_row)
            roles.append(["Junta Nueva Esperanza", "Presidencia", "Convocar asambleas", "Asamblea"])
            roles.append(["Junta Nueva Esperanza", "Tesorería", "Cobrar y llevar caja", "Directiva"])
            buf = io.BytesIO()
            wb.save(buf)
            payload = {"filename": "taller-22-oct.xlsx", "content_base64": base64.b64encode(buf.getvalue()).decode(), "workshop_id": wid}
            r = client.post("/advisory/import", headers=h1, json=payload)
            preview = r.json()
            check(r.status_code == 200 and preview["dry_run"] and preview["operators_created"] == ["Junta Nueva Esperanza"]
                  and preview["lists_saved"] == 1 and preview["forms_saved"] == 2 and not preview["errors"],
                  f"vista previa: 1 operador nuevo, 1 lista y 2 formularios, sin errores {preview['errors']}")
            check(not any(o["name"] == "Junta Nueva Esperanza" for o in client.get("/advisory/operators", headers=h1).json()),
                  "la vista previa no grabó nada")
            r = client.post("/advisory/import", headers=h1, json={**payload, "dry_run": False})
            check(r.status_code == 200 and not r.json()["dry_run"] and r.json()["forms_saved"] == 2, "carga real")
            nueva = next(o for o in client.get("/advisory/operators", headers=h1).json() if o["name"] == "Junta Nueva Esperanza")
            check(nueva["origin"] == "import" and nueva["profile"]["size"] == 60 and nueva["profile"]["color"] == "raw"
                  and nueva["profile"]["shape"] == "onsite" and "G1" in nueva["stages_in_workshops"],
                  "el operador cargado queda con su ficha, en el taller")
            csv_text = ("operador;formulario;campo;fila;valor;estimado\n"
                        "Junta El Rosal;MA-G1-PROFILE;parish;;Remigio Crespo;\n"
                        "Junta El Rosal;MA-G1-ACTIONS;problem;1;Falta de cloración;\n"
                        "Junta El Rosal;MA-G1-ACTIONS;first_action;1;Capacitar al operador;\n"
                        "Junta El Rosal;MA-G1-MAP;source;;Verde;\n")
            r = client.post("/advisory/import", headers=h1, json={"filename": "plano.csv", "dry_run": False,
                                                                 "content_base64": base64.b64encode(csv_text.encode()).decode()})
            res = r.json()
            check(r.status_code == 200 and res["forms_saved"] == 2 and res["lists_saved"] == 0
                  and any("Faltan 7 de 8" in e["message"] for e in res["errors"]),
                  "archivo plano: graba ficha y tabla; la lista incompleta se informa y no se graba")
            check(client.post("/advisory/import", headers=h1, json={"filename": "x.pdf", "content_base64": "AAAA"}).status_code == 422,
                  "formato no soportado -> 422")

            print("5. Tablero y mapa")
            d = client.get("/advisory/dashboard", headers=h1).json()
            check(d["kpis"]["operators"] == 3 and d["kpis"]["trained"] == 3 and d["kpis"]["workshops"] == 1
                  and d["kpis"]["users_covered"] == 225, "indicadores: 3 operadores capacitados, 225 usuarios")
            feats = {f["name"]: f for f in d["map"]["features"]}
            check(feats["Junta San Juan"]["size"] == 120 and feats["Junta San Juan"]["color"] == "potable"
                  and feats["Junta San Juan"]["shape"] == "sewer", "mapa: tamaño, color y forma desde la ficha")
            check([o["code"] for o in d["map"]["legend"]["color"]["options"]] == ["potable", "raw", "none"],
                  "leyenda del color desde el catálogo")
            g1s = next(s for s in d["stages"] if s["code"] == "G1")
            check(g1s["operators_trained"] == 3, "la G1 se vio con los 3 operadores")

            print("6. Traspaso a su propia organización")
            code = f"e2e-sanjuan-{SUFFIX}"
            r = client.post(f"/advisory/operators/{sj}/transfer", headers=h1,
                            json={"code": code, "admin_username": "presidencia", "admin_password": "clave-sanjuan"})
            check(r.status_code == 201 and r.json()["copied_runs"] == 1 and r.json()["copied_forms"] == 2,
                  f"organización creada con 1 lista y 2 formularios {r.text[:300]}")
            check(r.json()["findings_created"] == 1, "el rojo del mapa parlante es un hallazgo en la organización nueva")
            new_id = r.json()["organization"]["tenant_id"]
            created.append(new_id)
            hn = login(code, "presidencia", "clave-sanjuan")
            route = client.get("/process-route", headers=hn).json()
            g1n = {i["template_id"]: i["status"] for i in next(s for s in route["stages"] if s["code"] == "G1")["lists"]}
            check(g1n["MA-G1-MAP"] == "done" and g1n["MA-G1-PROFILE"] == "done" and g1n["MA-G1-PRIORITIES"] == "draft",
                  "el operador ve sus guías como las dejó el asesor")
            check(client.get("/forms/MA-G1-PROFILE", headers=hn).json()["data"]["fields"]["users"]["value"] == 120,
                  "su ficha trae los datos")
            r = client.put("/forms/MA-G1-PRIORITIES", headers=hn, json={"data": {"rows": [
                {"problem": "Fugas frecuentes en la red", "priority": "high"}]}})
            check(r.status_code == 200 and r.json()["progress"]["complete"], "el operador sigue llenando desde su organización")
            me = client.get("/auth/me", headers=hn).json()
            check(me["tenant_kind"] == "provider" and me["tenant_display_name"] == "Junta San Juan", "nace como prestador con su nombre")
            check(client.put(f"/advisory/operators/{sj}/forms/MA-G1-PROFILE", headers=h1,
                             json={"data": {"fields": {}}}).status_code == 409, "el registro de la ONG queda de solo lectura")
            check(client.post(f"/advisory/operators/{sj}/transfer", headers=h1,
                              json={"code": code + "x", "admin_username": "a", "admin_password": "clave-larga"}).status_code == 409,
                  "no se traspasa dos veces")
            check(client.get(f"/advisory/operators/{sj}", headers=h1).json()["transferred_tenant_name"] == code,
                  "la ONG ve a qué organización se traspasó")

            print("7. Aislamiento y permisos")
            r = client.post("/organizations", headers=hp, json={**body, "name": f"e2e-otra-ong-{SUFFIX}"})
            created.append(r.json()["tenant_id"])
            ho = login(f"e2e-otra-ong-{SUFFIX}", "admin", "clave-admin-apt")
            check(client.get("/advisory/operators", headers=ho).json() == [], "otra ONG no ve los operadores")
            check(client.get(f"/advisory/operators/{ops['Junta El Rosal']}", headers=ho).status_code == 404, "ni su detalle")
            check(client.get("/advisory/operators", headers=hj).status_code == 409, "un prestador no usa la asesoría")
            check(client.get("/process-route", headers=hj).json()["stages"][0]["lists"][0]["status"] == "never",
                  "la junta no ve nada de la ONG en sus guías")
        finally:
            for tid in reversed(created):
                with conn.transaction():
                    with tenant_scope(conn, tid):
                        for table in ("form_entry_version", "form_entry", "workshop_attendee", "workshop_operator",
                                      "checklist_answer", "checklist_run", "finding", "advised_operator", "workshop",
                                      "tenant_pack", "app_user"):
                            conn.execute(f"DELETE FROM {table} WHERE tenant_id = %s", (tid,))
                conn.execute("DELETE FROM tenant WHERE id = %s", (tid,))
            print("Limpieza: organizaciones de prueba y sus datos borrados")
    print("SPRINT E1 E2E OK")


if __name__ == "__main__":
    if len(sys.argv) != 2:
        print(__doc__)
        sys.exit(2)
    run(sys.argv[1])
