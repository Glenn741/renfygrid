"""Verificacion end-to-end real de Track F, sprints Q1 y Q2 (sectores y
menu como datos, fronteras, equipos y su hoja de vida; migraciones 0048 a
0050) -- HTTP real (FastAPI TestClient), JWT real, Postgres real con el rol
de aplicacion (RLS activo), sin mocks.

Que prueba, en espanol llano:
  1. La plataforma crea una comercializadora de energia; sin sector no se
     puede. Su menu sale de los datos: ve Energia, Mantenimiento y MDM, y
     nada de agua.
  2. El catalogo sale del paquete CO-CREG-038 (tipos de punto, frecuencias,
     fichas de la Circular 119). La ventana de "por vencer" es parametro.
  3. Frontera: alta, codigo repetido, datos obligatorios, tipo de punto
     sugerido por consumo.
  4. Equipos: alta, instalacion en su posicion (con los errores de tipo y de
     posicion ocupada), estado "sin registro" hasta calibrar.
  5. Bitacora: calibracion -> al dia con su vencimiento; retirar el sello de
     la tapa principal -> vencida; recalibrar -> al dia. El rol de aplicacion
     no puede editar ni borrar la bitacora.
  6. Cambio de medidor: retiro a laboratorio e instalacion del nuevo ->
     novedad "Cambio de medidor principal" pendiente de XM; marcarla
     reportada.
  7. Hoja de vida: una seccion de la Circular 119 se guarda y versiona;
     un campo que no existe se rechaza.
  8. Carga por archivo: plantilla desde el catalogo; vista previa sin
     grabar; carga real; volver a cargar no duplica.
  9. Permisos y aislamiento: un operador registra en la bitacora pero no
     crea fronteras; un acueducto no ve las fronteras ni el modulo.
 10. Mapa tematico (Q2.5): tipo de cliente, mercado y estado comercial con
     su catalogo; PQR con la taxonomia de la Superservicios (ligada por NIU,
     causales por sector); el mapa refleja cada dimension con su conteo; vistas
     guardadas; carga de PQR por archivo idempotente; permisos.
Al final borra SOLO las organizaciones de prueba y lo que crearon.

Uso:
    python verify_energy_assets_end_to_end.py "postgresql://renfygrid_app:renfygrid_app_dev_only@localhost:5455/renfygrid"

Contra un servicio ya corriendo:
    E2E_BASE_URL=http://10.12.200.27:8105 python verify_energy_assets_end_to_end.py "<dsn del servicio>"
"""

from __future__ import annotations

import base64
import io
import os
import sys
import uuid
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "common"))
sys.path.insert(0, str(Path(__file__).resolve().parent))

import psycopg  # noqa: E402
from psycopg.types.json import Json  # noqa: E402

from renmeter_common.db import tenant_scope  # noqa: E402
from renmeter_common.user_service import create_app_user  # noqa: E402
from verify_advisory_end_to_end import _HttpClient, check  # noqa: E402

JWT_SECRET = "e2e-f-secret"
ORDER_SIGNING_SECRET = "e2e-f-order-secret"
SUFFIX = uuid.uuid4().hex[:6]
CONFIG = {"session_ttl_seconds": 3600, "timezone": "America/Bogota", "currency": "COP", "locale": "es-CO"}


def run(dsn: str) -> None:
    from openpyxl import Workbook, load_workbook

    base = os.environ.get("E2E_BASE_URL")
    if base:
        client = _HttpClient(base)
    else:
        os.environ["RENFYGRID_DSN"] = dsn
        os.environ["RENFYGRID_JWT_SECRET"] = JWT_SECRET
        os.environ["RENFYGRID_ORDER_SIGNING_SECRET"] = ORDER_SIGNING_SECRET
        import main  # despues de fijar el entorno
        from fastapi.testclient import TestClient
        client = TestClient(main.app)
    created: list[str] = []

    def login(org: str, user: str, password: str) -> dict:
        r = client.post("/auth/login", json={"tenant_id": org, "email": user, "password": password})
        assert r.status_code == 200, r.text
        return {"Authorization": f"Bearer {r.json()['access_token']}"}

    with psycopg.connect(dsn, autocommit=True) as conn:
        with conn.cursor() as cur:
            for name, kind in ((f"e2e-plataforma-{SUFFIX}", "platform"), (f"e2e-acueducto-{SUFFIX}", "provider")):
                cur.execute("INSERT INTO tenant (name, config, kind) VALUES (%s, %s, %s) RETURNING id", (name, Json(CONFIG), kind))
                created.append(str(cur.fetchone()[0]))
        platform_id, water_id = created
        try:
            with conn.transaction():
                with tenant_scope(conn, water_id):
                    conn.execute("INSERT INTO tenant_sector (tenant_id, sector_code) VALUES (%s, 'water')", (water_id,))
            create_app_user(conn, platform_id, "root@plataforma.test", "clave-plataforma", "platform")
            create_app_user(conn, water_id, "admin@acueducto.test", "clave-acueducto", "supervisor")
            hp = login(f"e2e-plataforma-{SUFFIX}", "root@plataforma.test", "clave-plataforma")
            hw = login(f"e2e-acueducto-{SUFFIX}", "admin@acueducto.test", "clave-acueducto")

            print("1. Comercializadora de energía y su menú")
            org = f"e2e-energia-{SUFFIX}"
            body = {"name": org, "kind": "provider", "admin_username": "admin", "admin_password": "clave-admin-energia",
                    "display_name": "Comercializadora (prueba)", "packs": ["CO-CREG-038", "CO-SSPD"], "sectors": ["electricity"]}
            check(client.post("/organizations", headers=hp, json={**body, "sectors": []}).status_code == 422, "sin sector -> 422")
            r = client.post("/organizations", headers=hp, json=body)
            check(r.status_code == 201 and r.json()["sectors"] == ["electricity"], "la plataforma crea la comercializadora")
            org_id = r.json()["tenant_id"]
            created.append(org_id)
            ha = login(org, "admin", "clave-admin-energia")
            me = client.get("/auth/me", headers=ha).json()
            paths = [m["path"] for m in me["modules"]]
            check(me["sectors"] == ["electricity"], "la sesión dice que opera energía")
            check(paths[:4] == ["/energy", "/energy/map", "/energy/points", "/energy/equipment"], "el menú empieza por el módulo Energía (con el mapa)")
            check({"/maintenance", "/meters", "/vee", "/consumption", "/control", "/configuration"} <= set(paths),
                  "ve los módulos transversales (mantenimiento, MDM, configuración)")
            check(not {"/", "/quality", "/sanitation", "/inspections", "/network-balance", "/operations"} & set(paths),
                  "no ve nada de agua")
            water_paths = [m["path"] for m in client.get("/auth/me", headers=hw).json()["modules"]]
            check("/quality" in water_paths and "/energy" not in water_paths, "un acueducto ve agua y no energía")
            check([s["code"] for s in client.get("/sectors", headers=hp).json()] == ["electricity", "gas", "water"],
                  "catálogo de sectores")

            print("2. Catálogo del paquete y parámetros")
            cat = client.get("/assets/catalog", headers=ha).json()
            check(cat["pack_id"] == "CO-CREG-038" and [c["code"] for c in cat["point_classes"]] == ["1", "2", "3", "4", "5"],
                  "tipos de punto 1 a 5 del Código de Medida")
            freq = {(f["point_class"], f["equipment_class"], f["activity"]): f["every_months"] for f in cat["frequencies"]}
            check(freq[("1", "meter", "calibration")] == 24 and freq[("4", "meter", "calibration")] == 120
                  and freq[("2", "ct", "routine_test")] == 144, "frecuencias de la tabla 4 y del CNO 981")
            check(len(cat["forms"]) == 9 and sum(len(f["form_spec"]["fields"]) for f in cat["forms"]) > 100,
                  "las fichas de la Circular 119 salen del paquete")
            check(cat["due_soon_days"] is None, "la ventana de aviso no viene inventada")
            check(client.put("/settings/assets", headers=ha, json={"due_soon_days": 0}).status_code == 422, "ventana 0 -> 422")
            check(client.put("/settings/assets", headers=ha, json={"due_soon_days": 90}).json()["due_soon_days"] == 90,
                  "la organización fija la ventana de aviso")

            print("3. Frontera")
            point = {"code": f"FRT{SUFFIX}01", "name": "Planta Norte", "customer_name": "Industrias Prueba S.A.S.",
                     "city": "Barranquilla", "region": "Atlántico", "lat": 10.98, "lon": -74.8, "point_class": "2",
                     "voltage_kv": 13.2, "capacity_kva": 1500, "consumption_mwh_month": 620, "registered_on": "2022-03-01"}
            r = client.post("/metering-points", headers=ha, json=point)
            check(r.status_code == 201 and r.json()["suggested_point_class"] == "2", "frontera creada; el consumo sugiere tipo 2")
            pid = r.json()["id"]
            check(client.post("/metering-points", headers=ha, json=point).status_code == 409, "código repetido -> 409")
            check(client.post("/metering-points", headers=ha, json={**point, "code": "X1", "customer_name": ""}).status_code == 422,
                  "sin cliente -> 422")
            check(client.post("/metering-points", headers=ha, json={**point, "code": "X2", "point_class": "9"}).status_code == 422,
                  "tipo de punto que no existe -> 422")

            print("4. Equipos e instalación")
            def new_equipment(klass: str, serial: str) -> str:
                r = client.post("/equipment", headers=ha, json={"class_code": klass, "serial": serial, "brand": "Marca X",
                                                                "model": "M1", "year": 2021, "ownership": "rf"})
                assert r.status_code == 201, r.text
                return r.json()["id"]
            meter = new_equipment("meter", f"M{SUFFIX}A")
            cts = [new_equipment("ct", f"C{SUFFIX}{i}") for i in range(1, 4)]
            r = client.post(f"/equipment/{cts[0]}/install", headers=ha,
                            json={"point_id": pid, "position_code": "main_active", "installed_on": "2022-03-01"})
            check(r.status_code == 422, "un CT no va en la posición del medidor -> 422")
            r = client.post(f"/equipment/{meter}/install", headers=ha,
                            json={"point_id": pid, "position_code": "main_active", "installed_on": "2022-03-01"})
            check(r.status_code == 201 and r.json()["novelty"] is False, "medidor instalado (instalación inicial, sin novedad)")
            for i, ct in enumerate(cts, start=1):
                client.post(f"/equipment/{ct}/install", headers=ha,
                            json={"point_id": pid, "position_code": f"ct_{i}", "installed_on": "2022-03-01"})
            other = new_equipment("meter", f"M{SUFFIX}B")
            check(client.post(f"/equipment/{other}/install", headers=ha, json={"point_id": pid, "position_code": "main_active",
                                                                              "installed_on": "2022-03-02"}).status_code == 409,
                  "posición ocupada -> 409")
            hv = client.get(f"/metering-points/{pid}", headers=ha).json()
            check(len(hv["equipment"]) == 4 and hv["periodic_status"] == "missing", "4 equipos instalados, sin calibraciones -> sin registro")

            print("5. Bitácora: calibración, sello, recalibración")
            cal = {"certificate_number": "CAL-001", "laboratory": "Laboratorio Acreditado", "accredited": True,
                   "issued_on": "2023-06-01", "result": "pass"}
            r = client.post("/metering-events", headers=ha, json={"type_code": "calibration", "equipment_id": meter,
                                                                  "occurred_on": "2023-06-01",
                                                                  "data": {"fields": {k: {"value": v} for k, v in cal.items()}}})
            check(r.status_code == 201 and r.json()["point_id"] == pid, "calibración registrada; queda en la hoja de vida de la frontera")
            check(client.post("/metering-events", headers=ha, json={"type_code": "calibration", "equipment_id": meter,
                                                                    "occurred_on": "2023-06-01", "data": {"fields": {}}}).status_code == 422,
                  "calibración sin certificado -> 422")
            check(client.post("/metering-events", headers=ha, json={"type_code": "calibration", "equipment_id": meter,
                                                                    "occurred_on": "2099-01-01",
                                                                    "data": {"fields": {k: {"value": v} for k, v in cal.items()}}}).status_code == 422,
                  "fecha futura -> 422")
            for ct in cts:
                client.post("/metering-events", headers=ha, json={"type_code": "calibration", "equipment_id": ct, "occurred_on": "2021-12-01",
                                                                  "data": {"fields": {k: {"value": v} for k, v in {**cal, "issued_on": "2021-12-01"}.items()}}})
            eq = client.get(f"/equipment/{meter}", headers=ha).json()
            check(eq["periodic"]["status"] == "ok" and eq["periodic"]["due_on"] == "2027-06-01",
                  "medidor tipo 2: al día, vence a los 4 años (2027-06-01)")
            ct_eq = client.get(f"/equipment/{cts[0]}", headers=ha).json()
            check(ct_eq["periodic"]["activity"] == "routine_test" and ct_eq["periodic"]["due_on"] == "2033-12-01",
                  "CT: prueba de rutina a los 12 años desde su calibración inicial")
            r = client.post("/metering-events", headers=ha, json={"type_code": "seal_removed", "equipment_id": meter, "occurred_on": "2025-01-10",
                                                                  "data": {"fields": {"serial": {"value": "S-123"}, "main_cover": {"value": True}}}})
            check(r.status_code == 201, "retiro del sello de la tapa principal registrado")
            eq = client.get(f"/equipment/{meter}", headers=ha).json()
            check(eq["periodic"]["status"] == "overdue" and eq["periodic"].get("reason") == "invalidated",
                  "sin sello de tapa, el medidor debe recalibrarse (CREG 038 art. 11)")
            client.post("/metering-events", headers=ha, json={"type_code": "calibration", "equipment_id": meter, "occurred_on": "2025-01-20",
                                                              "data": {"fields": {k: {"value": v} for k, v in {**cal, "certificate_number": "CAL-002", "issued_on": "2025-01-20"}.items()}}})
            eq = client.get(f"/equipment/{meter}", headers=ha).json()
            check(eq["periodic"]["status"] == "ok" and eq["periodic"]["due_on"] == "2029-01-20", "recalibrado: al día de nuevo")
            with conn.transaction():
                with tenant_scope(conn, org_id):
                    try:
                        with conn.transaction():
                            conn.execute("UPDATE metering_event SET notes = 'x' WHERE tenant_id = %s", (org_id,))
                        edited = True
                    except psycopg.errors.InsufficientPrivilege:
                        edited = False
                    try:
                        with conn.transaction():
                            conn.execute("DELETE FROM metering_event WHERE tenant_id = %s", (org_id,))
                        deleted = True
                    except psycopg.errors.InsufficientPrivilege:
                        deleted = False
            check(not edited and not deleted, "la bitácora no se edita ni se borra con el rol de aplicación")

            print("6. Cambio de medidor y novedad para XM")
            r = client.post(f"/equipment/{meter}/remove", headers=ha, json={"removed_on": "2026-02-01", "reason": "Falla de display",
                                                                            "new_status": "lab"})
            check(r.status_code == 200 and client.get(f"/equipment/{meter}", headers=ha).json()["status"] == "lab",
                  "medidor retirado a laboratorio")
            r = client.post(f"/equipment/{other}/install", headers=ha,
                            json={"point_id": pid, "position_code": "main_active", "installed_on": "2026-02-01", "external_ref": "BF-0001"})
            check(r.status_code == 201 and r.json()["novelty"] is True, "medidor nuevo instalado: es un cambio")
            pending = client.get("/metering-events/xm-pending", headers=ha).json()
            change = [e for e in pending if e["type_code"] == "meter_change_main"]
            check(len(change) == 1 and change[0]["external_ref"] == "BF-0001", "novedad 'Cambio de medidor principal' pendiente de XM, con su orden de campo")
            hv = client.get(f"/metering-points/{pid}", headers=ha).json()
            check(any(h["removed_on"] == "2026-02-01" for h in hv["installation_history"]) and hv["xm_pending"] >= 1,
                  "la hoja de vida guarda el historial de instalaciones")
            r = client.post("/metering-events/xm-report", headers=ha, json={"event_ids": [change[0]["id"]], "reported_on": "2026-02-03"})
            check(r.json()["reported"] == 1 and not [e for e in client.get("/metering-events/xm-pending", headers=ha).json()
                                                     if e["type_code"] == "meter_change_main"], "novedad marcada como reportada")

            print("7. Secciones de la hoja de vida (Circular 119)")
            r = client.put(f"/metering-points/{pid}/sections/CO-HV-GENERAL", headers=ha,
                           json={"data": {"fields": {"c2_16": {"value": "Trifásica"}, "c2_17": {"value": "industrial"}}}})
            check(r.status_code == 200 and r.json()["version"] == 1 and r.json()["data"]["fields"]["c2_16"]["value"] == "trifasica",
                  "información general guardada (la opción se reconoce por su nombre)")
            r = client.put(f"/metering-points/{pid}/sections/CO-HV-GENERAL", headers=ha,
                           json={"data": {"fields": {"c2_16": {"value": "trifasica"}, "c2_99": {"value": "x"}}}})
            check(r.status_code == 422, "un campo que no está en la Circular -> 422")
            r = client.put(f"/metering-points/{pid}/sections/CO-HV-GENERAL", headers=ha,
                           json={"data": {"fields": {"c2_16": {"value": "trifasica"}}}})
            check(r.json()["version"] == 2, "cada guardado deja una versión")
            r = client.patch(f"/equipment/{other}", headers=ha, json={"attributes": {"fields": {"c3_9": {"value": "Estático"},
                                                                                                "c3_18": {"value": 0.5}}}})
            check(r.status_code == 200 and r.json()["attributes"]["fields"]["c3_9"]["value"] == "estatico",
                  "ficha técnica del medidor con los campos de la sección 3")

            print("8. Carga por archivo")
            r = client.get("/assets/import/template", headers=ha)
            wb = load_workbook(io.BytesIO(r.content))
            check(r.status_code == 200 and wb.sheetnames == ["Fronteras", "Equipos", "PQR", "Códigos"], "plantilla con fronteras, equipos, PQR y códigos")
            codes = [row[1] for row in wb["Códigos"].iter_rows(min_row=2, values_only=True) if row[0] == "Posición"]
            check("main_active" in codes and "ct_1" in codes, "los códigos de la plantilla salen del catálogo")
            wb = Workbook()
            ws = wb.active
            ws.title = "Fronteras"
            ws.append(["Código de la frontera (SIC)", "Nombre de la frontera", "Usuario (cliente)", "Tipo de punto", "Consumo mensual (MWh)"])
            ws.append([f"FRT{SUFFIX}02", "Bodega Sur", "Comercial Prueba", "Tipo 3", 80])
            ws.append([f"FRT{SUFFIX}03", "Local Centro", "Tienda Prueba", "4", 12])
            ws2 = wb.create_sheet("Equipos")
            ws2.append(["Tipo de equipo", "Número de serie", "Marca", "Frontera donde está instalado", "Posición", "Fecha de instalación",
                        "Fecha de la última calibración", "Certificado de calibración", "Laboratorio de calibración"])
            ws2.append(["Medidor", f"IM{SUFFIX}1", "Marca Y", f"FRT{SUFFIX}02", "main_active", "2024-05-01", "2024-04-20", "C-77", "Lab A"])
            ws2.append(["meter", f"IM{SUFFIX}2", "Marca Y", f"FRT{SUFFIX}03", "Medidor principal (energía activa)", "2019-01-10",
                        "2019-01-02", "C-78", "Lab A"])
            ws2.append(["meter", f"IM{SUFFIX}3", "Marca Y", "NO-EXISTE", "main_active", "2024-01-01", None, None, None])
            out = io.BytesIO()
            wb.save(out)
            payload = {"filename": "carga.xlsx", "content_base64": base64.b64encode(out.getvalue()).decode()}
            r = client.post("/assets/import", headers=ha, json=payload)
            res = r.json()
            check(r.status_code == 200 and res["dry_run"] and res["points_created"] == 2 and res["equipment_created"] == 2
                  and res["calibrations"] == 2 and len(res["errors"]) == 1 and res["errors"][0]["row"] == 4,
                  "vista previa: 2 fronteras, 2 equipos, 2 calibraciones y 1 fila con error")
            check(client.get("/metering-points", headers=ha).json()["total"] == 1, "la vista previa no grabó nada")
            res = client.post("/assets/import", headers=ha, json={**payload, "dry_run": False}).json()
            check(res["points_created"] == 2 and res["installed"] == 2, "carga real")
            listing = client.get("/metering-points", headers=ha).json()
            by_code = {p["code"]: p for p in listing["points"]}
            check(by_code[f"FRT{SUFFIX}02"]["periodic_status"] == "ok" and by_code[f"FRT{SUFFIX}03"]["periodic_status"] == "ok",
                  "tipo 3 calibrado en 2024 y tipo 4 calibrado en 2019: ambos al día")
            res = client.post("/assets/import", headers=ha, json={**payload, "dry_run": False}).json()
            check(res["points_updated"] == 2 and res["equipment_existing"] == 2 and res["calibrations"] == 0 and res["installed"] == 0,
                  "volver a cargar el mismo archivo no duplica nada")
            summary = client.get("/assets/summary", headers=ha).json()
            check(summary["points_total"] == 3 and summary["equipment_total"] == 7 and summary["points_by_class"]["2"] == 1,
                  "tablero: 3 fronteras y 7 equipos")

            print("9. Permisos y aislamiento")
            r = client.post("/users", headers=ha, json={"email": "tecnico", "password": "clave-tecnico", "role": "operator"})
            check(r.status_code == 201, "la comercializadora crea un técnico")
            ht = login(org, "tecnico", "clave-tecnico")
            r = client.post("/metering-events", headers=ht, json={"type_code": "verification_visit", "point_id": pid,
                                                                    "occurred_on": "2026-03-01",
                                                                    "data": {"fields": {"result": {"value": "ok"}}}})
            check(r.status_code == 201, f"el técnico registra una visita de verificación ({r.status_code} {r.text[:200]})")
            check(client.post("/metering-points", headers=ht, json={**point, "code": "X9"}).status_code == 403,
                  "pero no crea fronteras")
            check(client.get("/metering-points", headers=hw).json()["total"] == 0, "el acueducto no ve las fronteras")
            check(client.get(f"/metering-points/{pid}", headers=hw).status_code == 404, "ni su hoja de vida")

            print("10. Mapa temático y PQR (Q2.5)")
            r = client.patch(f"/metering-points/{pid}", headers=ha, json={"customer_class": "Industrial", "market": "non_regulated",
                                                                            "commercial_status": "arrears", "niu": f"NIU{SUFFIX}"})
            check(r.status_code == 200 and r.json()["customer_class"] == "industrial" and r.json()["commercial_status"] == "arrears",
                  "tipo de cliente, mercado y estado comercial (por nombre o por código)")
            check(client.patch(f"/metering-points/{pid}", headers=ha, json={"commercial_status": "inventado"}).status_code == 422,
                  "estado comercial fuera del catálogo -> 422")
            pcat = client.get("/customer-requests/catalog", headers=ha).json()
            codes = {c["code"] for c in pcat["causes"]}
            check({"101", "124", "200", "208"} <= codes and "207" not in codes and "130" not in codes,
                  "causales de la Superservicios que aplican a energía (sin las de acueducto ni aseo)")
            body_pqr = {"account": f"NIU{SUFFIX}", "filing_number": "RAD-E2E-1", "filed_on": "2026-09-01", "request_type": "1",
                        "cause_code": "101"}
            r = client.post("/customer-requests", headers=ha, json=body_pqr)
            check(r.status_code == 201 and r.json()["linked"], "PQR registrada y ligada a la frontera por su NIU")
            req_id = r.json()["id"]
            check(client.post("/customer-requests", headers=ha, json=body_pqr).status_code == 409, "misma PQR con la misma causal -> 409")
            check(client.post("/customer-requests", headers=ha, json={**body_pqr, "filing_number": "RAD-E2E-2", "cause_code": "207"}).status_code == 422,
                  "causal que no aplica a energía -> 422")
            r = client.post("/customer-requests", headers=ha, json={**body_pqr, "filing_number": "RAD-E2E-3", "cause_code": "200",
                                                                      "request_type": "2", "account": "NO-EXISTE"})
            check(r.status_code == 201 and not r.json()["linked"], "PQR de una cuenta desconocida queda sin frontera")
            check(len(client.get("/customer-requests", headers=ha).json()) == 2, "lista de PQR de la organización")
            m = client.get("/assets/map", headers=ha).json()
            idx = {d["key"]: i for i, d in enumerate(m["dimensions"])}
            check({"calibration", "comms", "commercial_status", "customer_class", "pqr_open", "pqr_cause", "meter_brand",
                   "modem_medium", "consumption_range"} <= set(idx), "el mapa trae las dimensiones del catálogo")
            mp = next(pt for pt in m["points"] if pt["id"] == pid)
            check(mp["v"][idx["pqr_open"]] == "yes" and mp["v"][idx["pqr_cause"]] == ["101"] and mp["v"][idx["pqr_group"]] == ["F"],
                  "la frontera aparece con PQR abierta de facturación, causal 101")
            check(mp["v"][idx["commercial_status"]] == "arrears" and mp["v"][idx["customer_class"]] == "industrial"
                  and mp["v"][idx["meter_brand"]] == "Marca X", "estado comercial, tipo de cliente y marca del medidor en el mapa")
            check(mp["v"][idx["comms"]] == "no_link", "sin medidor vinculado en el HES: 'sin telemedida vinculada'")
            cc = {c["code"]: c["count"] for c in m["dimensions"][idx["customer_class"]]["categories"]}
            check(cc.get("industrial") == 1, "las categorías traen su conteo")
            r = client.patch(f"/customer-requests/{req_id}", headers=ha, json={"status": "closed", "response_code": "1",
                                                                                 "responded_on": "2026-09-10"})
            check(r.status_code == 200, "PQR respondida y cerrada")
            m = client.get("/assets/map", headers=ha).json()
            mp = next(pt for pt in m["points"] if pt["id"] == pid)
            check(mp["v"][idx["pqr_open"]] == "no", "al cerrarla deja de contar como abierta en el mapa")
            view = {"paint": "commercial_status", "size": "size_consumption", "filters": {"customer_class": ["industrial"]},
                    "layers": ["calibration"]}
            r = client.post("/assets/map-views", headers=ha, json={"name": "Industriales en mora", "config": view})
            check(r.status_code == 201, "vista guardada")
            check(client.post("/assets/map-views", headers=ha, json={"name": "industriales en mora", "config": view}).status_code == 409,
                  "nombre de vista repetido -> 409")
            check(client.post("/assets/map-views", headers=ha, json={"name": "Mala", "config": {"x": 1}}).status_code == 422,
                  "configuración de vista inválida -> 422")
            views = client.get("/assets/map-views", headers=ha).json()
            check(len(views) == 1 and views[0]["config"]["filters"] == {"customer_class": ["industrial"]}, "la vista se lee igual")
            check(client.delete(f"/assets/map-views/{views[0]['id']}", headers=ha).status_code == 200, "vista borrada")
            wb = Workbook()
            ws = wb.active
            ws.title = "PQR"
            ws.append(["Radicado", "NIU o código de la frontera", "Fecha de radicación", "Tipo de trámite", "Causal (código de 3 dígitos)", "Estado"])
            ws.append(["RAD-E2E-10", f"FRT{SUFFIX}02", "2026-08-15", "Reclamación", "124", "Abierta"])
            ws.append(["RAD-E2E-11", f"FRT{SUFFIX}03", "2026-08-16", "Queja", 202, "open"])
            ws.append(["RAD-E2E-12", f"FRT{SUFFIX}03", "2026-08-16", "Queja", "999", "open"])
            out = io.BytesIO()
            wb.save(out)
            pqr_payload = {"filename": "pqr.xlsx", "content_base64": base64.b64encode(out.getvalue()).decode(), "dry_run": False}
            res = client.post("/assets/import", headers=ha, json=pqr_payload).json()
            check(res["requests_created"] == 2 and res["requests_unlinked"] == 0 and len(res["errors"]) == 1,
                  "carga de PQR por archivo: 2 ligadas por código de frontera y 1 causal inválida")
            res = client.post("/assets/import", headers=ha, json=pqr_payload).json()
            check(res["requests_created"] == 0 and res["requests_existing"] == 2, "volver a cargar las PQR no las duplica")
            check(client.post("/customer-requests", headers=ht, json={**body_pqr, "filing_number": "RAD-E2E-9"}).status_code == 403,
                  "el técnico no registra PQR (permiso requests.manage)")
            check(client.get("/assets/map", headers=hw).json()["total"] == 0, "el acueducto ve un mapa vacío")
        finally:
            for tid in reversed(created):
                with conn.transaction():
                    conn.execute("SELECT e2e_purge_metering_assets(%s)", (tid,))
                    with tenant_scope(conn, tid):
                        for table in ("tenant_pack", "tenant_sector", "tenant_module", "app_user"):
                            conn.execute(f"DELETE FROM {table} WHERE tenant_id = %s", (tid,))
                conn.execute("DELETE FROM tenant WHERE id = %s", (tid,))
            print("Limpieza: organizaciones de prueba y sus datos borrados")
    print("TRACK F Q1-Q2 E2E OK")


if __name__ == "__main__":
    if len(sys.argv) != 2:
        print(__doc__)
        sys.exit(2)
    run(sys.argv[1])
